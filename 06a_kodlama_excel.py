#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
06a_kodlama_excel.py - Kodlama dosyasını düzgün bir Excel dosyasına çevirir.

CSV dosyaları Türkçe Excel'de tek sütuna düşebiliyor. Bu betik aynı içeriği
xlsx olarak yazar: metin sütunu geniş ve satır kaydırmalı, kod sütununda
E / H / ? açılır menüsü, ikinci sayfada kod kitabı.

Kullanım:
  python 06a_kodlama_excel.py
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent
CODE_DIR = BASE / "veri" / "kodlama"
KAYNAK = CODE_DIR / "dogrulama_insan_alt.csv"
HEDEF = CODE_DIR / "dogrulama_insan_alt.xlsx"

KOD_KITABI = [
    ("SORU", "Bu mesaj, dışarıya dönük bir talep eyleminin parçası mı?"),
    ("", ""),
    ("E yaz", "Mesajın kendisi kampanya metni (etiketli talep, dilekçe/mektup metni)"),
    ("E yaz", "Eyleme çağırıyor (tweet atalım, imzalayalım, CİMER'e yazalım, destek olalım)"),
    ("E yaz", "Kampanyayı örgütlüyor (etkinlik planı, etiket önerisi, taktik uyarısı)"),
    ("E yaz", "Eyleme katılmak için soru soruyor (nasıl imzalanır, RT ne zaman atılır)"),
    ("", ""),
    ("H yaz", "Kampanya hakkında konuşuyor ama çağırmıyor (katılım azdı, işe yaramıyor)"),
    ("H yaz", "Durum sorusu (etkinlik ne zaman, devam ediyor mu)"),
    ("H yaz", "Panel / webinar / eğitim / canlı yayın duyurusu"),
    ("H yaz", "İş ilanı, tez anketi, haber paylaşımı, mevzuat bilgisi"),
    ("H yaz", "Sohbet (KPSS muhabbeti, puan sorusu, teşekkür, emoji, +1)"),
    ("H yaz", "Grup içi organizasyon (tercih grubu kurma, puan grubu açma)"),
    ("", ""),
    ("KURAL 1", "Eylem çağrısı yoksa H"),
    ("KURAL 2", "Mesaj tek başına anlaşılmıyorsa H (sadece link, '+', 'evet')"),
    ("KURAL 3", "Savunuculuk mesleki olmak zorunda değil (deprem yardımı, başka meslek kampanyası da E)"),
    ("KURAL 4", "Gerçekten karar veremezsen ? yaz, hesaplamadan çıkarılır"),
]


def main():
    if not KAYNAK.exists():
        raise SystemExit(f"{KAYNAK} bulunamadı. Önce 05_dogrulama_ornegi.py çalıştırın.")
    df = pd.read_csv(KAYNAK)
    if "kod" not in df.columns:
        df["kod"] = ""
    df["kod"] = df["kod"].fillna("")

    with pd.ExcelWriter(HEDEF, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="kodlama", index=False)
        pd.DataFrame(KOD_KITABI, columns=["durum", "açıklama"]).to_excel(
            w, sheet_name="kod_kitabi", index=False)

        ws = w.book["kodlama"]
        genislik = {"sira": 7, "tarih": 12, "metin": 110, "kod": 10}
        for i, ad in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = genislik.get(ad, 16)
        sutunlar = list(df.columns)
        kod_sutun = get_column_letter(sutunlar.index("kod") + 1)
        metin_sutun = (get_column_letter(sutunlar.index("metin") + 1)
                       if "metin" in sutunlar else None)

        for satir in ws.iter_rows(min_row=2, max_row=len(df) + 1):
            for hucre in satir:
                hucre.alignment = Alignment(wrap_text=True, vertical="top")
        for hucre in ws[1]:
            hucre.font = Font(bold=True)
            hucre.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.freeze_panes = "A2"

        dv = DataValidation(type="list", formula1='"E,H,?"', allow_blank=True,
                            showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{kod_sutun}2:{kod_sutun}{len(df) + 1}")

        kk = w.book["kod_kitabi"]
        kk.column_dimensions["A"].width = 12
        kk.column_dimensions["B"].width = 95
        for satir in kk.iter_rows(min_row=1):
            for hucre in satir:
                hucre.alignment = Alignment(wrap_text=True, vertical="top")

    print(f"Hazır: {HEDEF}")
    print(f"{len(df)} satır. 'kod' sütununda E / H / ? seçeneklerini kullan.")
    print("Kod kitabı ikinci sayfada. Bitince 06_dogrulama_hesapla.py çalıştır.")
    if metin_sutun:
        print(f"Metin sütunu dar gelirse: {metin_sutun} sütununu genişlet ya da hücreye "
              "tıklayıp üstteki formül çubuğundan oku.")


if __name__ == "__main__":
    main()
