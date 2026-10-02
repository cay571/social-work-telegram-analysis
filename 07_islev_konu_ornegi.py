#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
07_islev_konu_ornegi.py - İşlev ve konu kodlaması için örneklem hazırlar.

Üretilen dosyalar (veri/kodlama/):
  islev_konu_ornegi.csv    -> 600 mesaj (bana göndereceğin dosya)
  islev_konu_insan.xlsx    -> bunların 70'i, senin kodlaman için (açılır menülü)
  islev_konu_anahtar.csv   -> sira-msg_id-yil eşlemesi (kodlarken açılmaz)

Örneklem, tekrarsız metin kümesinden yıllara göre tabakalı çekilir.

Kullanım:
  python 07_islev_konu_ornegi.py
  python 07_islev_konu_ornegi.py --n 600 --insan 70
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent
MESAJLAR = BASE / "veri" / "islenmis" / "mesajlar.parquet"
CODE_DIR = BASE / "veri" / "kodlama"
TOHUM = 20260919
METIN_SINIR = 500

ISLEV = [
    ("bilgi_talebi", "Soru soruyor, bilgi ya da yardım istiyor"),
    ("bilgi_verme", "Soruya yanıt veriyor, bilgi/açıklama sunuyor, yönlendiriyor"),
    ("duyuru", "İlan, mevzuat, haber, sonuç, link paylaşımı (talep ya da yanıt değil)"),
    ("mobilizasyon", "Eylem çağrısı ya da kampanya metni (tweet atalım, imzalayalım, dilekçe metni)"),
    ("degerlendirme", "Yorum, eleştiri, tartışma, görüş bildirme, tahmin yürütme"),
    ("duygu_destek", "Dert yanma, moral verme, tebrik, dayanışma, şikâyet-duygu ifadesi"),
    ("sohbet", "Selam, teşekkür, şaka, emoji, tek kelimelik onay ('evet', '+1')"),
    ("diger", "Grup işleyişi, üyelik, reklam, anket, anlaşılmayan metin"),
]
KONU = [
    ("istihdam", "Atama, kadro, alım, mülakat, sözleşmeli/4B, yerleştirme"),
    ("sinav_puan", "KPSS hazırlık, puan, sıralama, tercih süreci"),
    ("egitim_akademik", "Lisans/lisansüstü, AÖF, ders, tez, staj, akademik yaşam"),
    ("uygulama_vaka", "Mesleki uygulama, vaka, danışan, kurumda iş yapış biçimi"),
    ("ozluk_mevzuat", "Unvan, özlük hakları, maaş, tayin, mevzuat, sendika"),
    ("toplumsal_olay", "Deprem/afet, gündem olayları, toplumsal sorunlar"),
    ("grup_ici", "Grubun kendi işleyişi, üyelik, kurallar, grup kurma"),
    ("belirsiz", "Konu anlaşılmıyor ya da yukarıdakilerin dışında"),
]


def tabakali(kaynak, n, tohum):
    if len(kaynak) <= n:
        return kaynak.copy()
    paylar = (kaynak.groupby("yil").size() / len(kaynak) * n).round().astype(int)
    parcalar = [kaynak[kaynak["yil"] == y].sample(min(k, int((kaynak["yil"] == y).sum())),
                                                  random_state=tohum)
                for y, k in paylar.items() if k > 0]
    ornek = pd.concat(parcalar)
    return ornek.sample(min(n, len(ornek)), random_state=tohum)


def insan_dosyasi(sayfa, yol):
    with pd.ExcelWriter(yol, engine="openpyxl") as w:
        sayfa.to_excel(w, sheet_name="kodlama", index=False)
        kitap = pd.DataFrame(
            [("İŞLEV", "Mesaj ne YAPIYOR? Tek kod seç.")] +
            [("islev", f"{k}: {a}") for k, a in ISLEV] +
            [("", ""), ("KONU", "Mesaj ne HAKKINDA? Tek kod seç.")] +
            [("konu", f"{k}: {a}") for k, a in KONU] +
            [("", ""),
             ("KURAL 1", "Birden fazla işlev varsa, mesajın ağırlık merkezini seç"),
             ("KURAL 2", "Soru işareti varsa ama soru değilse (retorik) bilgi_talebi değildir"),
             ("KURAL 3", "Bir soruya cevap veriyorsa bilgi_verme, kendi başına paylaşımsa duyuru"),
             ("KURAL 4", "Kampanya hakkında konuşuyor ama çağırmıyorsa degerlendirme (mobilizasyon değil)"),
             ("KURAL 5", "Karar veremezsen ilgili sütuna ? yaz")],
            columns=["alan", "açıklama"])
        kitap.to_excel(w, sheet_name="kod_kitabi", index=False)

        ws = w.book["kodlama"]
        genislik = {"sira": 7, "tarih": 12, "metin": 100, "islev": 18, "konu": 18}
        for i, ad in enumerate(sayfa.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = genislik.get(ad, 16)
        for satir in ws.iter_rows(min_row=2, max_row=len(sayfa) + 1):
            for hucre in satir:
                hucre.alignment = Alignment(wrap_text=True, vertical="top")
        for hucre in ws[1]:
            hucre.font = Font(bold=True)
            hucre.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.freeze_panes = "A2"

        for ad, kodlar in (("islev", ISLEV), ("konu", KONU)):
            harf = get_column_letter(list(sayfa.columns).index(ad) + 1)
            dv = DataValidation(type="list", allow_blank=True, showDropDown=False,
                                formula1='"' + ",".join([k for k, _ in kodlar] + ["?"]) + '"')
            ws.add_data_validation(dv)
            dv.add(f"{harf}2:{harf}{len(sayfa) + 1}")

        kk = w.book["kod_kitabi"]
        kk.column_dimensions["A"].width = 14
        kk.column_dimensions["B"].width = 95
        for satir in kk.iter_rows(min_row=1):
            for hucre in satir:
                hucre.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=600)
    ap.add_argument("--insan", type=int, default=70)
    args = ap.parse_args()

    CODE_DIR.mkdir(parents=True, exist_ok=True)
    m = pd.read_parquet(MESAJLAR)
    havuz = m[m["metin_analizi"] & ~m["is_dup_same_sender"] & (m["n_token"] >= 1)]
    ornek = tabakali(havuz, args.n, TOHUM).sample(frac=1, random_state=TOHUM).reset_index(drop=True)
    ornek["sira"] = range(1, len(ornek) + 1)

    kodlanacak = pd.DataFrame({
        "sira": ornek["sira"],
        "tarih": ornek["date_tr"].dt.strftime("%Y-%m-%d"),
        "metin": ornek["text_light"].str.slice(0, METIN_SINIR).str.replace(r"\s+", " ", regex=True),
    })
    kodlanacak.to_csv(CODE_DIR / "islev_konu_ornegi.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame({"sira": ornek["sira"], "msg_id": ornek["msg_id"], "yil": ornek["yil"]}).to_csv(
        CODE_DIR / "islev_konu_anahtar.csv", index=False, encoding="utf-8-sig")

    alt = kodlanacak.sample(min(args.insan, len(kodlanacak)), random_state=TOHUM + 5).sort_values("sira")
    alt = alt.assign(islev="", konu="")
    insan_dosyasi(alt, CODE_DIR / "islev_konu_insan.xlsx")

    print(f"Havuz (tekrarsız metin kümesi): {len(havuz):,}")
    print(f"Örneklem: {len(ornek)} mesaj | senin kodlayacağın alt küme: {len(alt)}")
    print("\nYıllara göre dağılım:")
    print(ornek.groupby("yil").size().to_string())
    print(f"\nDosyalar {CODE_DIR}:")
    print("  islev_konu_ornegi.csv   -> bana gönder")
    print("  islev_konu_insan.xlsx   -> islev ve konu sütunlarını sen doldur")
    print("  islev_konu_anahtar.csv  -> açma, hesaplamada kullanılacak")


if __name__ == "__main__":
    main()
